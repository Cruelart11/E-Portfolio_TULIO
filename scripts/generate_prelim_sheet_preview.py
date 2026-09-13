from html import escape
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT / "public" / "documents" / "prelim" / "pt-p1-workbook.xlsx"
OUTPUT_PATH = ROOT / "public" / "documents" / "prelim" / "pt-p1-workbook-preview.html"


def display_value(cell):
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.4f}".rstrip("0").rstrip(".")
    return str(value)


workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
sheet = workbook.active

merged_starts = {}
merged_covered = set()
for merged_range in sheet.merged_cells.ranges:
    min_col, min_row, max_col, max_row = range_boundaries(str(merged_range))
    merged_starts[(min_row, min_col)] = (max_row - min_row + 1, max_col - min_col + 1)
    for row in range(min_row, max_row + 1):
        for column in range(min_col, max_col + 1):
            if (row, column) != (min_row, min_col):
                merged_covered.add((row, column))

rows = []
for row_index in range(1, sheet.max_row + 1):
    cells = [f'<th scope="row">{row_index}</th>']
    for column_index in range(1, sheet.max_column + 1):
        if (row_index, column_index) in merged_covered:
            continue

        cell = sheet.cell(row=row_index, column=column_index)
        rowspan, colspan = merged_starts.get((row_index, column_index), (1, 1))
        attributes = []
        if rowspan > 1:
            attributes.append(f'rowspan="{rowspan}"')
        if colspan > 1:
            attributes.append(f'colspan="{colspan}"')
        if cell.font.bold:
            attributes.append('class="is-bold"')
        attribute_text = f" {' '.join(attributes)}" if attributes else ""
        value = escape(display_value(cell)).replace("\n", "<br>")
        cells.append(f"<td{attribute_text}>{value}</td>")
    rows.append(f"<tr>{''.join(cells)}</tr>")

column_headers = "".join(
    f'<th scope="col">{get_column_letter(column)}</th>'
    for column in range(1, sheet.max_column + 1)
)

html = f"""<!doctype html>
<html lang="en">
  <head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>PT-P1 Workbook Preview</title>
    <style>
      :root {{ color-scheme: dark; font-family: Arial, sans-serif; }}
      * {{ box-sizing: border-box; }}
      body {{ margin: 0; color: #e6edf3; background: #0d1117; }}
      table {{ min-width: 1900px; border-collapse: collapse; font-size: 12px; }}
      th, td {{ min-width: 108px; max-width: 280px; padding: 8px 10px; border: 1px solid #30363d; vertical-align: top; white-space: normal; }}
      thead th, tbody th {{ position: sticky; color: #8b98ac; background: #161b22; font-family: monospace; font-weight: 400; }}
      thead th {{ top: 0; z-index: 2; }}
      tbody th {{ left: 0; min-width: 44px; z-index: 1; text-align: right; }}
      thead th:first-child {{ left: 0; z-index: 3; min-width: 44px; }}
      td {{ background: #0d1117; line-height: 1.45; }}
      tr:nth-child(even) td {{ background: #10151d; }}
      td.is-bold {{ color: #7ee787; font-weight: 700; }}
    </style>
  </head>
  <body>
    <table aria-label="Exercise PT-P1 workbook data">
      <thead><tr><th aria-label="Row number"></th>{column_headers}</tr></thead>
      <tbody>{''.join(rows)}</tbody>
    </table>
  </body>
</html>
"""

OUTPUT_PATH.write_text(html, encoding="utf-8")
print(f"Generated {OUTPUT_PATH} from {sheet.max_row} rows and {sheet.max_column} columns.")
